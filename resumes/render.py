"""
resumes/render.py — one-column .docx / .txt output. Plain, boring,
ATS-parseable: standard headings, real bullet lists, no tables, no text boxes,
no columns.
"""

from docx import Document
from docx.shared import Pt

from .ats import resume_to_text


def to_txt(m: dict) -> str:
    return resume_to_text(m)


def to_docx(m: dict) -> bytes:
    import io
    doc = Document()
    for st in doc.sections:
        st.top_margin = st.bottom_margin = Pt(48)
        st.left_margin = st.right_margin = Pt(64)

    b = m.get("basics") or {}
    h = doc.add_heading(level=0)
    r = h.add_run(str(b.get("name") or "Your Name"))
    r.font.size = Pt(20)
    contact = "  |  ".join(str(x) for x in (b.get("email"), b.get("phone"), b.get("location")) if x)
    if contact:
        p = doc.add_paragraph()
        p.add_run(contact).font.size = Pt(10)

    def section(title: str):
        doc.add_heading(title, level=1)

    if b.get("summary"):
        section("SUMMARY")
        doc.add_paragraph(str(b["summary"]).strip())

    if m.get("skills"):
        section("SKILLS")
        doc.add_paragraph(", ".join(m["skills"]))

    if m.get("work"):
        section("EXPERIENCE")
        for w in m["work"]:
            when = " - ".join(str(x) for x in (w.get("start"), w.get("end")) if x)
            head = f"{w.get('role', '')}" + (f"  |  {w['company']}" if w.get("company") else "")
            if when:
                head += f"  ({when})"
            p = doc.add_paragraph()
            p.add_run(head).bold = True
            for bl in w.get("bullets") or []:
                doc.add_paragraph(str(bl), style="List Bullet")

    if m.get("education"):
        section("EDUCATION")
        for e in m["education"]:
            extra = " - ".join(str(x) for x in (e.get("degree"), e.get("year")) if x)
            line = str(e.get("school") or "") + (f"  -  {extra}" if extra else "")
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
