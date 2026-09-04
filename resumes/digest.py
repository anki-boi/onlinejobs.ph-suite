"""Digest multiple resume source files into a fact corpus for the LLM CV builder.

The corpus is deterministic raw material (text, bullets, frequent terms) —
the LLM decides what to keep, but it may only use what appears here.
"""
import re
from pathlib import Path

import docx
import pymupdf

TEXT_EXTS = {".txt", ".md"}
OFFICE_EXTS = {".pdf", ".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}

BULLET_RE = re.compile(r"^\s*[•●◦▪\-\*]\s*(.+)$")
STOP = set(
    "the and for with you your that this from have has had are was were be been "
    "will can able using use used into over under more most other some all any "
    "each their there here what which when how who whom then them they she he his "
    "her its our ours not no yes per via etc if or but also one two three years "
    "year work worked working team teams role roles job jobs position positions "
    "company companies business service services customer customers client clients "
    "data information report reports system systems process processes task tasks "
    "detail details high low new good best better ensure ensured ensure accurate "
    "accurate ability abilities strong good great excellent successful success "
    "results result outcomes outcome outcome quality efficient efficiently "
    "effectively effective ability skill skills"
)


def read_file(path: Path) -> str:
    """Extract text from a pdf/docx/plain file. Returns '' on failure."""
    try:
        if path.suffix.lower() == ".pdf":
            doc = pymupdf.open(str(path))
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
        elif path.suffix.lower() == ".docx":
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
        return re.sub(r"[ \t]+", " ", text).strip()
    except Exception:
        return ""


def extract_bullets(text: str) -> list[str]:
    """All bullet lines, verbatim, deduplicated (order-preserving)."""
    seen, out = set(), []
    for line in text.splitlines():
        m = BULLET_RE.match(line)
        if m:
            b = m.group(1).strip()
            if 8 < len(b) < 300 and b.lower() not in seen:
                seen.add(b.lower())
                out.append(b)
    return out


def extract_terms(text: str, top_n: int = 60) -> list[str]:
    """Most frequent alphanumeric tokens outside the stoplist — a skill hint
    list for the LLM to filter, not a skill list itself."""
    freq: dict[str, int] = {}
    for tok in re.findall(r"[A-Za-z][A-Za-z0-9+#./-]{2,}", text.lower()):
        if tok in STOP:
            continue
        freq[tok] = freq.get(tok, 0) + 1
    return [t for t, _ in sorted(freq.items(), key=lambda kv: -kv[1])[:top_n]]


def collect_files(sources) -> list[Path]:
    """Expand config sources (files or directories) into an ordered file list."""
    files = []
    for src in sources or []:
        p = Path(src)
        if p.is_dir():
            files.extend(sorted(c for c in p.rglob("*") if c.is_file()))
        elif p.is_file():
            files.append(p)
    # masters.json first if present (identity anchor), then the rest
    def order(p: Path):
        return (0 if p.name.lower() in ("masters.json", "master.json") else 1, str(p))
    return sorted(files, key=order)


def digest(sources, per_source_cap: int = 6000, bullet_cap: int = 120,
           term_cap: int = 60, total_cap: int = 20000) -> dict:
    """Digest files from the given sources into a JSON-serializable corpus."""
    bullets, terms_text, source_entries, images = [], [], [], []
    total = 0
    for f in collect_files(sources):
        ext = f.suffix.lower()
        if ext in IMAGE_EXTS:
            images.append(str(f))
            continue
        if ext not in TEXT_EXTS | OFFICE_EXTS and ext != ".json":
            continue
        if ext == ".json":
            text = f.read_text(encoding="utf-8", errors="replace")
        else:
            text = read_file(f)
        if not text:
            continue
        bullets.extend(extract_bullets(text))
        terms_text.append(text)
        cap = per_source_cap
        entry = {"file": str(f), "chars": len(text), "text": text[:cap]}
        source_entries.append(entry)
        total += len(text)
        if total >= total_cap:
            break
    seen, uniq_bullets = set(), []
    for b in bullets:
        if b.lower() not in seen:
            seen.add(b.lower())
            uniq_bullets.append(b)
    terms = extract_terms("\n".join(terms_text), top_n=term_cap)
    return {
        "sources": source_entries,
        "bullets": uniq_bullets[:bullet_cap],
        "terms": terms[:term_cap],
        "images": images,
        "chars": total,
    }


def corpus_text(corpus: dict, cap: int = 16000) -> str:
    """Render the digest as one prompt-sized text block."""
    parts = []
    for s in corpus["sources"]:
        parts.append(f"=== SOURCE: {s['file']} ({s['chars']} chars) ===\n{s['text']}")
    if corpus["bullets"]:
        parts.append("=== ALL BULLET POINTS (verbatim material) ===\n"
                     + "\n".join("- " + b for b in corpus["bullets"]))
    if corpus["terms"]:
        parts.append("=== FREQUENT TERMS (skill hints, filter these) ===\n"
                     + ", ".join(corpus["terms"]))
    if corpus["images"]:
        parts.append("=== NON-TEXT FILES (mention as attachments if relevant) ===\n"
                     + "\n".join(corpus["images"]))
    text = "\n\n".join(parts)
    return text[:cap]
